"""
Script to seed the Supabase database with location hierarchy data (Location -> Blocks -> Floors -> Rooms)
from Excel and CSV files in the project root.
"""

import os
import sys
import csv
import openpyxl

# Add project root to python path to resolve app imports
sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from dotenv import load_dotenv
load_dotenv()

from app.utils.db import get_db

def seed_hierarchy():
    admin = get_db()
    
    print("Initializing Seeding process...")

    # 1. Create or retrieve the root Location "UIET"
    print("Upserting Location: UIET...")
    loc_res = admin.table("locations").upsert({"name": "UIET"}, on_conflict="name").execute()
    if not loc_res.data:
        raise Exception("Failed to upsert location 'UIET'")
    location_id = loc_res.data[0]["id"]
    print(f"Location 'UIET' ID: {location_id}")

    # Define the block files mapping
    files_blocks = [
        {"file": "UIET BLOCK 1.xlsx", "block": "Block 1", "type": "xlsx"},
        {"file": "UIET BLOCK 2.xlsx", "block": "Block 2", "type": "xlsx"},
        {"file": "UIET BLOCK 3.xlsx", "block": "Block 3", "type": "xlsx"},
        {"file": "UIET WORKSHOP.csv", "block": "Workshop", "type": "csv"},
        {"file": "CANTEEN1.xlsx", "block": "Canteen", "type": "xlsx"},
    ]

    for fb in files_blocks:
        filename = fb["file"]
        block_name = fb["block"]
        file_type = fb["type"]
        filepath = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), filename)

        if not os.path.exists(filepath):
            print(f"Warning: File {filename} not found at {filepath}, skipping...")
            continue

        print(f"\nProcessing Block: '{block_name}' from file {filename}...")

        # Upsert Block
        block_res = admin.table("blocks").upsert({
            "location_id": location_id,
            "name": block_name
        }, on_conflict="location_id,name").execute()
        if not block_res.data:
            print(f"Error: Failed to upsert block '{block_name}', skipping file.")
            continue
        block_id = block_res.data[0]["id"]
        print(f"Block '{block_name}' ID: {block_id}")

        # Read rooms and floors from file
        rows = []
        if file_type == "xlsx":
            try:
                wb = openpyxl.load_workbook(filepath)
                sheet = wb.active
                for r in range(2, sheet.max_row + 1):
                    floor_val = sheet.cell(row=r, column=1).value
                    room_num_val = sheet.cell(row=r, column=2).value
                    room_name_val = sheet.cell(row=r, column=3).value
                    if floor_val is not None:
                        rows.append((floor_val, room_num_val, room_name_val))
            except Exception as e:
                print(f"Error reading xlsx {filename}: {e}")
                continue
        else: # csv
            try:
                with open(filepath, mode='r', encoding='utf-8') as f:
                    reader = csv.reader(f)
                    next(reader, None) # skip header
                    for r in reader:
                        if len(r) >= 1 and r[0]:
                            floor_val = r[0]
                            room_num_val = r[1] if len(r) > 1 else None
                            room_name_val = r[2] if len(r) > 2 else None
                            rows.append((floor_val, room_num_val, room_name_val))
            except Exception as e:
                print(f"Error reading csv {filename}: {e}")
                continue

        print(f"Found {len(rows)} raw rows to parse.")

        # Extract unique floors and insert them first to cache IDs
        unique_floors = set()
        for floor_val, _, _ in rows:
            floor_name = str(floor_val).strip()
            # Normalize floor name
            if floor_name.lower() == "ground":
                floor_name = "Ground Floor"
            elif floor_name.lower() == "first":
                floor_name = "First Floor"
            unique_floors.add(floor_name)

        floor_cache = {}
        for floor_name in sorted(unique_floors):
            floor_res = admin.table("floors").upsert({
                "block_id": block_id,
                "name": floor_name
            }, on_conflict="block_id,name").execute()
            if floor_res.data:
                floor_cache[floor_name] = floor_res.data[0]["id"]
                print(f"  Floor '{floor_name}' ID: {floor_res.data[0]['id']}")
            else:
                print(f"  Error: Failed to upsert floor '{floor_name}'")

        # Batch insert rooms
        rooms_to_insert = []
        for floor_val, room_num_val, room_name_val in rows:
            floor_name = str(floor_val).strip()
            if floor_name.lower() == "ground":
                floor_name = "Ground Floor"
            elif floor_name.lower() == "first":
                floor_name = "First Floor"

            floor_id = floor_cache.get(floor_name)
            if not floor_id:
                continue

            # Handle room number
            if room_num_val is None or str(room_num_val).strip() == "" or str(room_num_val).strip().upper() == "N/A":
                room_number = "N/A"
            else:
                room_number = str(room_num_val).strip()

            # Handle room name
            room_name = str(room_name_val).strip() if room_name_val is not None else ""

            rooms_to_insert.append({
                "floor_id": floor_id,
                "room_number": room_number,
                "name": room_name
            })

        # Insert rooms in batches of 100
        batch_size = 100
        total_inserted = 0
        for i in range(0, len(rooms_to_insert), batch_size):
            batch = rooms_to_insert[i:i + batch_size]
            try:
                res = admin.table("rooms").upsert(batch, on_conflict="floor_id,room_number,name").execute()
                total_inserted += len(res.data)
            except Exception as e:
                print(f"Error bulk upserting rooms: {e}")
                # Fallback to row-by-row in case of validation issues
                for room in batch:
                    try:
                        admin.table("rooms").upsert(room, on_conflict="floor_id,room_number,name").execute()
                        total_inserted += 1
                    except Exception as e_row:
                        print(f"Failed to insert single room {room}: {e_row}")

        print(f"Successfully upserted {total_inserted} rooms for block '{block_name}'.")

    print("\nLocation hierarchy seeding process completed!")

if __name__ == "__main__":
    seed_hierarchy()
